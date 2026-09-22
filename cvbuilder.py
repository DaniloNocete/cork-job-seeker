"""Tailored-CV builder: master CV in -> job-specific CV out (html/txt/docx)."""
import re
import datetime
from cvmatch import extract_terms, STOPWORDS


# ---------- import: plain text -> structured draft ----------

DATE_RANGE = (r"((?:[A-Za-z]{3,9}\.?\s+)?(?:19|20)\d{2}\s*(?:[-–—]|\sto\s)\s*"
              r"(?:(?:[A-Za-z]{3,9}\.?\s+)?(?:19|20)\d{2}|present|current|now|date|ongoing)"
              r"|(?:19|20)\d{2}\s*[-–—]+\s*(?:19|20)\d{2})")

SECTION_ALIASES = {
    "experience": ["work experience", "employment", "employment history", "work history",
                   "professional experience", "experience", "additional experience", "other experience",
                   "hospitality experience"],
    "education": ["education", "qualifications", "academic", "training"],
    "skills": ["skills", "key skills", "core skills", "technical skills", "competencies"],
    "summary": ["profile", "personal profile", "summary", "professional summary",
                "personal statement", "objective", "about me"],
}


def parse_cv_text(raw):
    """Best-effort parse of a pasted plain-text CV into structured sections."""
    out = {"summary": "", "skills": "", "availability": "", "experience": [], "education": []}
    if not raw or not raw.strip():
        return out
    lines = [l.rstrip() for l in raw.replace("\r", "").split("\n")]
    sections, current = {}, "head"
    buf = []

    def flush():
        if current in sections:
            sections[current].extend(buf)
        else:
            sections[current] = list(buf)
        buf.clear()

    for line in lines:
        low = line.strip().lower().rstrip(":")
        matched_sec = None
        if low and len(low) < 40:
            for sec, aliases in SECTION_ALIASES.items():
                if low in aliases:
                    matched_sec = sec
                    break
        if matched_sec:
            flush()
            current = matched_sec
        else:
            buf.append(line)
    flush()

    def nonempty(ls):
        return [l.strip() for l in ls if l.strip()]

    out["summary"] = " ".join(nonempty(sections.get("summary", sections.get("head", []))))[:800]
    skill_lines = [l for l in sections.get("skills", []) if not re.search(r"availab", l, re.I)]
    out["skills"] = ", ".join(nonempty(skill_lines))
    # availability line anywhere
    for l in lines:
        if re.search(r"availab", l, re.I):
            out["availability"] = l.strip()[:160]
            break
    # experience: treat non-empty lines as entries; lines starting with - or • are bullets
    entries, cur = [], None
    for l in nonempty(sections.get("experience", [])):
        if re.match(r"^[-•*·]\s*", l):
            if cur:
                cur["bullets"].append(re.sub(r"^[-•*·]\s*", "", l))
        else:
            cur = {"role": l, "org": "", "dates": "", "bullets": []}
            m = re.search(DATE_RANGE, l, re.I)
            if m:
                cur["dates"] = m.group(0).strip()
                cur["role"] = re.sub(r"[,|]?\s*" + re.escape(m.group(0)), "", l).strip(" ,|")
            entries.append(cur)
    if not entries:
        entries = [{"role": l, "org": "", "dates": "", "bullets": []} for l in nonempty(sections.get("experience", []))]
    out["experience"] = entries
    edu, cur = [], None
    for l in nonempty(sections.get("education", [])):
        if re.match(r"^[-•*·]\s*", l) and cur:
            cur["award"] += " | " + re.sub(r"^[-•*·]\s*", "", l)
        else:
            cur = {"award": l, "org": "", "dates": ""}
            m = re.search(DATE_RANGE, l, re.I)
            if m:
                cur["dates"] = m.group(0).strip()
                cur["award"] = re.sub(r"[,|]?\s*" + re.escape(m.group(0)), "", l).strip(" ,|")
            edu.append(cur)
    out["education"] = edu
    return out


# ---------- tailoring ----------

def _term_hits(text, terms):
    low = (text or "").lower()
    return [t for t in terms if re.search(r"(?<![a-z])" + re.escape(t) + r"(?![a-z])", low)]


def tailor_cv(profile, master, job, analysis):
    """Build a tailored CV dict from master CV + job + match analysis."""
    job_text = job.get("description", "")
    job_skills, job_extra = extract_terms(job_text)
    all_terms = list(dict.fromkeys(job_skills + job_extra))

    name = profile.get("name") or "[Your Name]"
    matched = analysis.get("matched", [])

    # -- summary: role-targeted opener + kept sentences containing job keywords
    opener = (f"Motivated and reliable {profile.get('headline') or 'candidate'} "
              f"applying for the {job.get('title', 'advertised')} position at "
              f"{job.get('company', 'your company')}.")
    if matched:
        opener += f" Brings hands-on strengths in {', '.join(matched[:4])}."
    if master.get("availability"):
        opener += f" {master['availability'].rstrip('.')}."
    kept = []
    for sent in re.split(r"(?<=[.!?])\s+", master.get("summary", "")):
        if sent.strip() and _term_hits(sent, all_terms) and sent.strip() not in opener:
            kept.append(sent.strip())
    summary = " ".join([opener] + kept[:2])

    # -- skills: matched first (ordered by job relevance), then the rest
    my_skills = [s.strip() for s in re.split(r"[,;\n]", master.get("skills", "")) if s.strip()]
    def rel(s):
        hits = len(_term_hits(s, all_terms))
        return (-hits, s.lower())
    skills_matched = sorted([s for s in my_skills if _term_hits(s, all_terms)], key=rel)
    skills_other = [s for s in my_skills if s not in skills_matched]

    # -- experience: reorder entries & bullets by keyword hits (content unchanged)
    exp = []
    for e in master.get("experience", []):
        body = " ".join([e.get("role", ""), e.get("org", ""), " ".join(e.get("bullets", []))])
        score = len(_term_hits(body, all_terms))
        bullets = sorted(e.get("bullets", []), key=lambda b: -len(_term_hits(b, all_terms)))
        exp.append({**e, "bullets": bullets, "_score": score})
    exp.sort(key=lambda e: -e["_score"])
    for e in exp:
        e.pop("_score", None)

    return {
        "name": name,
        "contact": " · ".join(x for x in [profile.get("phone"), profile.get("email"),
                                          profile.get("location") or "Cork, Ireland"] if x),
        "summary": summary,
        "skills_matched": skills_matched,
        "skills_other": skills_other,
        "experience": exp,
        "education": master.get("education", []),
        "availability": master.get("availability", ""),
        "suggestions": [k for k in analysis.get("missing", []) if k not in " ".join(my_skills).lower()][:8],
        "job_title": job.get("title", ""),
        "company": job.get("company", ""),
        "generated": datetime.date.today().isoformat(),
    }


# ---------- renderers ----------

def render_txt(cv):
    L = []
    L.append(cv["name"].upper())
    L.append(cv["contact"])
    bar = "=" * 60
    L += [bar, "PROFILE", cv["summary"], ""]
    if cv["skills_matched"]:
        L += [bar, "KEY SKILLS FOR THIS ROLE", ", ".join(cv["skills_matched"]), ""]
    if cv["skills_other"]:
        L += [bar, "OTHER SKILLS", ", ".join(cv["skills_other"]), ""]
    if cv["experience"]:
        L += [bar, "EXPERIENCE"]
        for e in cv["experience"]:
            head = " | ".join(x for x in [e.get("role"), e.get("org")] if x)
            dates = f" ({e['dates']})" if e.get("dates") else ""
            L.append(f"{head}{dates}")
            L += [f"  • {b}" for b in e.get("bullets", [])]
            L.append("")
    if cv["education"]:
        L += [bar, "EDUCATION"]
        for e in cv["education"]:
            head = " | ".join(x for x in [e.get("award"), e.get("org")] if x)
            dates = f" ({e['dates']})" if e.get("dates") else ""
            L.append(f"{head}{dates}")
        L.append("")
    return "\n".join(L).strip()


def render_html(cv):
    from html import escape as esc
    def sec(t): return f"<h4 style='margin:14px 0 6px;font-size:12px;letter-spacing:1.2px;color:#4f8cff;text-transform:uppercase;border-bottom:1px solid #2a3550;padding-bottom:4px'>{esc(t)}</h4>"
    parts = [f"<div style='text-align:center;margin-bottom:10px'>"
             f"<div style='font-size:20px;font-weight:800'>{esc(cv['name'])}</div>"
             f"<div style='font-size:12.5px;color:#93a0b8'>{esc(cv['contact'])}</div></div>"]
    parts.append(sec("Profile") + f"<p style='font-size:13.5px'>{esc(cv['summary'])}</p>")
    if cv["skills_matched"]:
        parts.append(sec("Key skills for this role") +
                     "".join(f"<span class='kw good'>{esc(s)}</span>" for s in cv["skills_matched"]))
    if cv["skills_other"]:
        parts.append(sec("Other skills") +
                     "".join(f"<span class='kw'>{esc(s)}</span>" for s in cv["skills_other"]))
    if cv["experience"]:
        parts.append(sec("Experience"))
        for e in cv["experience"]:
            head = " | ".join(esc(x) for x in [e.get("role"), e.get("org")] if x)
            dates = f" <span style='color:#93a0b8'>({esc(e.get('dates',''))})</span>" if e.get("dates") else ""
            bl = "".join(f"<li style='font-size:13px;margin:2px 0'>{esc(b)}</li>" for b in e.get("bullets", []))
            parts.append(f"<div style='margin-bottom:8px'><b style='font-size:13.5px'>{head}</b>{dates}"
                         + (f"<ul style='margin:4px 0 0 18px'>{bl}</ul>" if bl else "") + "</div>")
    if cv["education"]:
        parts.append(sec("Education"))
        for e in cv["education"]:
            head = " | ".join(esc(x) for x in [e.get("award"), e.get("org")] if x)
            dates = f" <span style='color:#93a0b8'>({esc(e.get('dates',''))})</span>" if e.get("dates") else ""
            parts.append(f"<div style='font-size:13.5px'><b>{head}</b>{dates}</div>")
    return "".join(parts)


def render_docx(cv, path):
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.6)
        s.left_margin = s.right_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run(cv["name"])
    r.bold = True
    r.font.size = Pt(18)
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = c.add_run(cv["contact"])
    cr.font.size = Pt(9.5)
    cr.font.color.rgb = RGBColor(0x60, 0x6A, 0x80)

    def heading(text):
        p = doc.add_paragraph()
        p.space_before = Pt(10)
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x2B, 0x5C, 0xB8)
        pPr = p._p.get_or_add_pPr()
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "2"); bottom.set(qn("w:color"), "B8C4D8")
        pBdr.append(bottom); pPr.append(pBdr)
        return p

    heading("Profile")
    doc.add_paragraph(cv["summary"])
    if cv["skills_matched"]:
        heading("Key Skills for This Role")
        doc.add_paragraph(", ".join(cv["skills_matched"]))
    if cv["skills_other"]:
        heading("Other Skills")
        doc.add_paragraph(", ".join(cv["skills_other"]))
    if cv["experience"]:
        heading("Experience")
        for e in cv["experience"]:
            p = doc.add_paragraph()
            run = p.add_run(" | ".join(x for x in [e.get("role"), e.get("org")] if x))
            run.bold = True
            if e.get("dates"):
                dr = p.add_run(f"  ({e['dates']})")
                dr.font.color.rgb = RGBColor(0x60, 0x6A, 0x80)
                dr.font.size = Pt(9.5)
            for b in e.get("bullets", []):
                bp = doc.add_paragraph(b, style="List Bullet")
                bp.paragraph_format.space_after = Pt(2)
    if cv["education"]:
        heading("Education")
        for e in cv["education"]:
            p = doc.add_paragraph()
            run = p.add_run(" | ".join(x for x in [e.get("award"), e.get("org")] if x))
            run.bold = True
            if e.get("dates"):
                dr = p.add_run(f"  ({e['dates']})")
                dr.font.color.rgb = RGBColor(0x60, 0x6A, 0x80)
                dr.font.size = Pt(9.5)
    doc.save(path)
    return path
